"""Import TLC book-of-business (BOB) Excel sheets into empty TLC policy shells."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook

from .tlc_carriers import ensure_tlc_carrier
from .tlc_models import TLCPolicy, TLCPolicyTimelineEvent, TLCPremiumBreakdown


class BobSheetParseError(ValueError):
    """Raised when a BOB workbook cannot be read or has no usable rows."""


_HEADER_ALIASES = {
    "account name": "account_name",
    "named insured": "account_name",
    "insured": "account_name",
    "policy type": "policy_type",
    "line of business": "line_of_business",
    "lob": "line_of_business",
    "master company": "carrier",
    "company": "carrier",
    "carrier": "carrier",
    "policy number": "policy_number",
    "policy #": "policy_number",
    "term": "term",
    "effective date": "effective_date",
    "eff date": "effective_date",
    "annualized premium": "premium",
    "premium": "premium",
    "written premium": "premium",
}


def _normalize_header(value) -> str:
    text = re.sub(r"\s+", " ", str(value or "").strip().lower())
    return text


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _parse_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m-%d-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_decimal(value) -> Decimal:
    if value is None or value == "":
        return Decimal("0.00")
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value)).quantize(Decimal("0.01"))
        except (InvalidOperation, ValueError):
            return Decimal("0.00")
    text = re.sub(r"[,$]", "", str(value).strip())
    try:
        return Decimal(text).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return Decimal("0.00")


def _term_months(value) -> int:
    text = str(value or "").strip().lower()
    match = re.search(r"(\d+)\s*month", text)
    if match:
        return max(1, int(match.group(1)))
    match = re.search(r"(\d+)\s*year", text)
    if match:
        return max(1, int(match.group(1)) * 12)
    digits = re.search(r"(\d+)", text)
    if digits:
        months = int(digits.group(1))
        return months if months > 0 else 12
    return 12


def _add_months(start: date, months: int) -> date:
    year = start.year + (start.month - 1 + months) // 12
    month = (start.month - 1 + months) % 12 + 1
    day = min(start.day, [31, 29 if year % 4 == 0 and (year % 100 != 0 or year % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][month - 1])
    return date(year, month, day)


def _clean_policy_number(value) -> str:
    return re.sub(r"\s+", "", _cell_str(value))


def _clean_carrier(value) -> str:
    text = _cell_str(value)
    text = re.sub(r"\s+APPLICATION\s*$", "", text, flags=re.I).strip()
    text = re.sub(r"\s+", " ", text)
    return text


@dataclass
class BobImportResult:
    created: int = 0
    skipped_duplicates: int = 0
    skipped_invalid: int = 0
    created_numbers: list[str] | None = None

    def __post_init__(self):
        if self.created_numbers is None:
            self.created_numbers = []


def parse_bob_workbook(upload) -> list[dict]:
    """Return normalized row dicts from the Detail sheet (or first sheet)."""
    data = upload.read() if hasattr(upload, "read") else upload
    if hasattr(upload, "seek"):
        try:
            upload.seek(0)
        except Exception:
            pass
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise BobSheetParseError(f"Could not read Excel file: {exc}") from exc

    sheet_name = "Detail" if "Detail" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise BobSheetParseError("The workbook is empty.")

    header_cells = [_normalize_header(c) for c in rows[0]]
    mapping = {}
    for idx, header in enumerate(header_cells):
        key = _HEADER_ALIASES.get(header)
        if key and key not in mapping:
            mapping[key] = idx

    required = ("account_name", "policy_number", "carrier")
    missing = [k for k in required if k not in mapping]
    if missing:
        raise BobSheetParseError(
            "Missing required columns: "
            + ", ".join(missing)
            + ". Expected Account Name, Policy Number, and Master Company."
        )

    parsed = []
    for raw in rows[1:]:
        if not raw or all(c is None or str(c).strip() == "" for c in raw):
            continue

        def col(key, default=""):
            idx = mapping.get(key)
            if idx is None or idx >= len(raw):
                return default
            return raw[idx]

        policy_number = _clean_policy_number(col("policy_number"))
        account_name = _cell_str(col("account_name"))
        carrier = _clean_carrier(col("carrier"))
        if not policy_number and not account_name:
            continue
        parsed.append(
            {
                "account_name": account_name,
                "policy_type_label": _cell_str(col("policy_type")),
                "line_of_business": _cell_str(col("line_of_business")),
                "carrier": carrier,
                "policy_number": policy_number,
                "term": _cell_str(col("term")),
                "effective_date": _parse_date(col("effective_date")),
                "premium": _parse_decimal(col("premium")),
            }
        )
    if not parsed:
        raise BobSheetParseError("No policy rows found in the sheet.")
    return parsed


def import_bob_rows_to_tlc(*, organization, space, rows, user=None) -> BobImportResult:
    """Create empty TLC policy shells from BOB rows (update later via Dec pages)."""
    result = BobImportResult()
    existing = set(
        TLCPolicy.objects.filter(organization=organization).values_list(
            "policy_number", flat=True
        )
    )
    # Case-insensitive duplicate check against cleaned numbers already in org.
    existing_upper = {str(n).strip().upper() for n in existing}

    for row in rows:
        policy_number = row["policy_number"]
        if not policy_number:
            result.skipped_invalid += 1
            continue
        if policy_number.upper() in existing_upper:
            result.skipped_duplicates += 1
            continue

        carrier = row["carrier"]
        if carrier:
            ensure_tlc_carrier(organization, carrier)

        effective = row["effective_date"]
        months = _term_months(row["term"])
        expiration = _add_months(effective, months) if effective else None
        form_of_business = row["line_of_business"] or row["policy_type_label"]
        notes_bits = [
            "Imported from BOB Excel sheet as an empty policy shell.",
            "Update from a declaration page PDF to fill vehicles, drivers, and schedule.",
        ]
        if row["policy_type_label"]:
            notes_bits.append(f"Sheet policy type: {row['policy_type_label']}.")
        if row["term"]:
            notes_bits.append(f"Term: {row['term']}.")

        policy = TLCPolicy.objects.create(
            organization=organization,
            space=space,
            policy_number=policy_number,
            carrier=carrier,
            named_insured=row["account_name"],
            form_of_business=form_of_business[:80],
            status=TLCPolicy.Status.PENDING,
            effective_date=effective,
            expiration_date=expiration,
            renewal_date=expiration,
            notes=" ".join(notes_bits),
            added_by=user,
        )
        TLCPremiumBreakdown.objects.create(
            policy=policy,
            total_written_premium=row["premium"],
        )
        policy.save()
        TLCPolicyTimelineEvent.objects.create(
            policy=policy,
            event_type=TLCPolicyTimelineEvent.EventType.QUOTE,
            title="Imported from BOB sheet",
            description=(
                f"Shell policy created for {row['account_name'] or policy_number}. "
                "Awaiting DEC page update."
            ),
            event_date=effective or date.today(),
            created_by=user,
        )
        existing_upper.add(policy_number.upper())
        result.created += 1
        result.created_numbers.append(policy_number)

    return result
