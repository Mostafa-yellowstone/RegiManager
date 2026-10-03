"""Join Kintone CSV exports into one profile per unique column.

An existing profile is never copied. If the unique number, driver license,
or the same name and phone is already in the office, that profile is skipped
and only missing vehicles, transactions, and documents are added.
"""

import csv
import hashlib
import io
import os
import re
import zipfile
from datetime import datetime
from decimal import Decimal, InvalidOperation

from django.core.files.base import ContentFile
from django.db import transaction
from django.utils import timezone
from django.utils.timezone import is_naive, make_aware
from openpyxl import load_workbook

from .models import Client, ClientNote, ServiceDocument, ServiceRecord, Vehicle

LINK_ALIASES = (
    "unique",
    "unique_id",
    "unique_key",
    "unique_column",
    "client_key",
    "client_number",
    "customer_number",
    "customer_id",
    "external_key",
    "link_key",
    "id_number",
)

FILE_RE = re.compile(r"\.(pdf|jpe?g|png|webp|heic|tiff?|docx?|xlsx?)$", re.I)
DATE_FORMATS = (
    "%b %d, %Y %I:%M %p",
    "%b %d, %Y %I:%M:%S %p",
    "%B %d, %Y %I:%M %p",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%m/%d/%Y %I:%M %p",
    "%m/%d/%Y",
    "%Y-%m-%d",
    "%m-%d-%Y",
)
SERVICE_TYPES = {key for key, _label in ServiceRecord.SERVICE_TYPES}


def header_key(value):
    text = str(value or "").strip().lower().replace("$", "")
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def cell_str(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%b %d, %Y %I:%M %p")
    return str(value).strip()


def pick(row, *names):
    for name in names:
        value = cell_str(row.get(header_key(name)))
        if value:
            return value
    return ""


def normalize_key(value):
    return cell_str(value).strip()


def alnum(value):
    return re.sub(r"[^A-Za-z0-9]", "", cell_str(value)).upper()


def phone_digits(value):
    return "".join(ch for ch in cell_str(value) if ch.isdigit())


def tidy_name(part):
    part = cell_str(part)
    if part.isupper():
        return part.title()
    return part


def split_last_first_middle(full):
    parts = [tidy_name(part) for part in re.sub(r"\s+", " ", cell_str(full)).split(" ") if part]
    if not parts:
        return "", "", ""
    if len(parts) == 1:
        return parts[0], "Client", ""
    return parts[1], parts[0], " ".join(parts[2:])


def person_from_row(row):
    first = tidy_name(pick(row, "first_name", "first", "firstname"))
    last = tidy_name(pick(row, "last_name", "last", "lastname"))
    middle = tidy_name(pick(row, "middle_name", "middle", "middlename"))
    if first and last:
        return first, last, middle
    return split_last_first_middle(pick(row, "name", "full_name", "client_name", "customer_name", "client"))


def parse_when(value):
    if isinstance(value, datetime):
        return value
    text = cell_str(value)
    if not text:
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def parse_year(value):
    digits = re.sub(r"\D", "", cell_str(value))
    if len(digits) == 4 and digits.startswith(("19", "20")):
        return int(digits)
    return None


def parse_money(value):
    text = cell_str(value).replace("$", "").replace(",", "")
    if not text:
        return Decimal("0.00")
    try:
        return Decimal(text)
    except InvalidOperation:
        return Decimal("0.00")


def split_filenames(value):
    names = []
    for part in re.split(r"[\n,;|]+", cell_str(value)):
        name = os.path.basename(part.strip().strip('"'))
        if name and FILE_RE.search(name):
            names.append(name)
    return names


def filenames_in_row(row):
    found = []
    file_headers = {
        "file",
        "filename",
        "file_name",
        "attachment",
        "attachments",
        "document",
        "documents",
        "document_name",
    }
    for key, value in row.items():
        if key in file_headers or FILE_RE.search(cell_str(value)):
            found.extend(split_filenames(value))
    return found


def rows_from_upload(uploaded):
    if uploaded is None:
        return []
    name = (getattr(uploaded, "name", "") or "").lower()
    uploaded.seek(0)
    if name.endswith(".xlsx"):
        workbook = load_workbook(filename=uploaded, read_only=True, data_only=True)
        sheet = workbook.active
        grid = list(sheet.iter_rows(values_only=True))
        workbook.close()
        if not grid:
            return []
        headers = [header_key(cell) for cell in grid[0]]
        rows = []
        for raw in grid[1:]:
            row = {}
            for idx, key in enumerate(headers):
                if key and idx < len(raw):
                    row[key] = raw[idx]
            if any(cell_str(value) for value in row.values()):
                rows.append(row)
        return rows

    raw_bytes = uploaded.read()
    try:
        text = raw_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw_bytes.decode("latin-1")
    sample = text[:4096]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    rows = []
    for raw in reader:
        row = {header_key(key): value for key, value in raw.items() if key}
        if any(cell_str(value) for value in row.values()):
            rows.append(row)
    return rows


def resolve_link_header(rows, requested):
    if not rows:
        raise ValueError("The clients file has no data rows.")
    headers = set(rows[0].keys())
    requested_key = header_key(requested)
    if requested_key:
        if requested_key in headers:
            return requested_key
        contains = [header for header in headers if requested_key in header]
        if len(contains) == 1:
            return contains[0]
        raise ValueError(
            f"Could not find a column named '{requested}'. Headers in the clients file: {', '.join(sorted(headers))}."
        )
    for alias in LINK_ALIASES:
        if alias in headers:
            return alias
    fuzzy = [header for header in headers if "unique" in header]
    if len(fuzzy) == 1:
        return fuzzy[0]
    raise ValueError(
        "Add the shared unique column to the clients export, or type its column name. "
        f"Headers found: {', '.join(sorted(headers))}."
    )


def require_link_header(rows, link_header, label):
    if not rows:
        return
    if link_header not in rows[0]:
        headers = ", ".join(sorted(rows[0].keys()))
        raise ValueError(
            f"The {label} file is missing the unique column '{link_header}'. "
            f"Use the same column on every export. Headers found: {headers}."
        )


def document_type_for(filename):
    lower = filename.lower()
    if "registration" in lower:
        return "registration"
    if "title" in lower:
        return "title"
    if "insurance" in lower:
        return "insurance_id"
    if "license" in lower:
        return "driver_license"
    if "mv82" in lower:
        return "mv82"
    return "other"


def service_type_for(label):
    text = cell_str(label).lower().replace("-", " ")
    key = re.sub(r"[^a-z0-9]+", "_", text).strip("_")
    if key in SERVICE_TYPES:
        return key
    if "motorcycle" in text and "reg" in text:
        return "motorcycle_registration"
    if "renew" in text:
        return "registration_renewal"
    if "title" in text:
        return "title_only"
    if "regist" in text:
        return "vehicle_registration"
    return "other"


class KintoneImporter:
    def __init__(self, organization, actor, link_header):
        self.organization = organization
        self.actor = actor
        self.link_header = link_header
        self.clients = {}
        self.files_for_key = {}
        self.motorcycle_keys = set()
        self.created_clients = 0
        self.skipped_clients = 0
        self.seen_client_ids = set()
        self.created_vehicles = 0
        self.skipped_vehicles = 0
        self.created_transactions = 0
        self.skipped_transactions = 0
        self.created_documents = 0
        self.skipped_documents = 0
        self.errors = []
        self._zip = None
        self._zip_index = {}
        self._license_index = None

    def remember(self, key, client, created):
        self.clients[key] = client
        if client.pk in self.seen_client_ids:
            return
        self.seen_client_ids.add(client.pk)
        if created:
            self.created_clients += 1
        else:
            self.skipped_clients += 1

    def license_index(self):
        if self._license_index is None:
            self._license_index = {}
            for candidate in Client.objects.filter(organization=self.organization).exclude(driver_license=""):
                self._license_index.setdefault(alnum(candidate.driver_license), candidate)
        return self._license_index

    def find_existing(self, key, row):
        found = Client.objects.filter(organization=self.organization, external_key=key).first()
        if found:
            return found
        license_values = {alnum(key)}
        license_number = pick(row, "driver_license", "dl", "dl_number", "license_number")
        if license_number:
            license_values.add(alnum(license_number))
        license_values.discard("")
        for value in license_values:
            found = self.license_index().get(value)
            if found is not None:
                return found
        phone = phone_digits(pick(row, "phone_number", "phone", "phone_no", "mobile", "tel"))
        first, last, _middle = person_from_row(row)
        if len(phone) >= 10 and first and last:
            candidates = Client.objects.filter(
                organization=self.organization,
                first_name__iexact=first,
                last_name__iexact=last,
            )
            for candidate in candidates:
                stored = phone_digits(candidate.phone_number)
                if len(stored) >= 10 and stored[-10:] == phone[-10:]:
                    return candidate
        return None

    def client_for(self, key, row):
        key = normalize_key(key)
        if not key:
            return None
        cached = self.clients.get(key)
        if cached is not None:
            return cached
        existing = self.find_existing(key, row)
        if existing is not None:
            if not existing.external_key:
                existing.external_key = key[:64]
                existing.save(update_fields=["external_key"])
            self.remember(key, existing, created=False)
            return existing

        first, last, middle = person_from_row(row)
        if not first or not last:
            self.errors.append(f"No name for unique number {key}, so that profile was skipped.")
            return None
        gender = pick(row, "gender", "sex").lower()
        if gender.startswith("m"):
            gender = "male"
        elif gender.startswith("f"):
            gender = "female"
        else:
            gender = ""
        state = pick(row, "state").upper()[:2]
        if state not in {code for code, _label in Client.US_STATES}:
            state = "NY"
        category = pick(row, "category", "account_type", "client_type").lower()
        phone = phone_digits(pick(row, "phone_number", "phone", "phone_no", "mobile", "tel"))
        license_number = normalize_key(pick(row, "driver_license", "dl", "dl_number", "license_number")) or key
        created_at = parse_when(pick(row, "created_datetime", "created_at", "date_created", "datetime", "date"))
        client = Client(
            organization=self.organization,
            first_name=first[:100],
            last_name=last[:100],
            middle_name=middle[:100],
            external_key=key[:64],
            driver_license=license_number[:50],
            phone_number=phone[:20],
            email=pick(row, "email", "email_address") or None,
            gender=gender or None,
            city=tidy_name(pick(row, "city"))[:100],
            state=state,
            zip_code=pick(row, "zip_code", "zip", "zipcode")[:10],
            county=tidy_name(pick(row, "county"))[:100],
            street_address=pick(row, "street_address", "street", "address")[:200],
            is_commercial=category in {"commercial", "business", "dealer"},
            source="walk_in",
        )
        if created_at is not None and is_naive(created_at):
            created_at = make_aware(created_at)
        with transaction.atomic():
            client.save()
            if created_at is not None:
                Client.objects.filter(pk=client.pk).update(created_at=created_at)
                client.created_at = created_at
        self.license_index()[alnum(client.driver_license)] = client
        self.remember(key, client, created=True)
        return client

    def note_files(self, key, row):
        names = filenames_in_row(row)
        bucket = self.files_for_key.setdefault(key, [])
        vin = pick(row, "vin", "vin_number").upper()
        plate = pick(row, "plate_number", "plate", "license_plate").upper()
        for name in names:
            bucket.append({"name": name, "vin": vin, "plate": plate})
            if "motorcycle" in name.lower():
                self.motorcycle_keys.add(key)

    def vehicle_for(self, client, row, create=True):
        vin = pick(row, "vin", "vin_number").upper()
        plate = pick(row, "plate_number", "plate", "license_plate").upper()
        record_number = pick(row, "record_number", "record_no", "record_id")
        if vin:
            other = (
                Vehicle.objects.filter(vin__iexact=vin, client__organization=self.organization)
                .exclude(client=client)
                .first()
            )
            if other:
                if create:
                    self.errors.append(f"VIN {vin} already belongs to {other.client.name}, so it was skipped.")
                    self.skipped_vehicles += 1
                return None
            current = Vehicle.objects.filter(client=client, vin__iexact=vin).first()
            if current:
                if create:
                    self.skipped_vehicles += 1
                return current
        elif plate:
            current = Vehicle.objects.filter(client=client, plate_number__iexact=plate).first()
            if current:
                if create:
                    self.skipped_vehicles += 1
                return current
        if not create:
            return None
        if not vin:
            vin = f"KT{re.sub(r'[^A-Z0-9]', '', (record_number or plate or client.external_key).upper())}"[:50]
            current = Vehicle.objects.filter(client=client, vin__iexact=vin).first()
            if current:
                self.skipped_vehicles += 1
                return current
        explicit_type = pick(row, "vehicle_type", "body_type").lower()
        if "motor" in explicit_type:
            vehicle_type = "motorcycle"
        elif explicit_type in {code for code, _label in Vehicle.VEHICLE_TYPES}:
            vehicle_type = explicit_type
        elif client.external_key in self.motorcycle_keys:
            vehicle_type = "motorcycle"
        else:
            vehicle_type = "passenger"
        vehicle = Vehicle(
            client=client,
            vin=vin[:50],
            is_legacy_vin=len(vin) != 17 or any(char in vin for char in "IOQ"),
            plate_number=plate[:50],
            year=parse_year(pick(row, "year", "model_year")),
            make=pick(row, "make", "brand")[:100],
            model=pick(row, "model")[:100],
            vehicle_type=vehicle_type,
            plate_type="motorcycle" if vehicle_type == "motorcycle" else "personal",
            vehicle_number=(f"KT-{record_number}" if record_number else "")[:50],
        )
        with transaction.atomic():
            vehicle.save()
        self.created_vehicles += 1
        return vehicle

    def add_transaction(self, client, row):
        label = pick(row, "service", "service_type", "transaction", "transaction_type", "description")
        if not label and not pick(row, "fee", "amount", "service_fee", "total", "price"):
            return
        service_type = service_type_for(label)
        when = parse_when(pick(row, "transaction_date", "date", "created_datetime", "datetime"))
        fee = parse_money(pick(row, "service_fee", "fee", "amount", "total", "price"))
        record_number = pick(row, "record_number", "record_no", "record_id")
        digest = hashlib.sha1(
            f"{self.organization.id}|{client.external_key}|{service_type}|{when}|{fee}|{record_number}".encode()
        ).hexdigest()[:16]
        case_id = f"KT{digest}"
        if ServiceRecord.objects.filter(case_id=case_id).exists():
            self.skipped_transactions += 1
            return
        vehicle = self.vehicle_for(client, row, create=True)
        if vehicle is None:
            return
        notes = "Imported from Kintone."
        if label and service_type == "other":
            notes = f"{notes} Service: {label}."
        if record_number:
            notes = f"{notes} Record {record_number}."
        record = ServiceRecord(
            organization=self.organization,
            handled_by=self.actor,
            vehicle=vehicle,
            service_type=service_type,
            service_fee=fee,
            transaction_date=when.date() if when else timezone.now().date(),
            notes=notes,
            case_id=case_id,
            client_name=client.name,
        )
        with transaction.atomic():
            record.save()
        self.created_transactions += 1

    def attach_files(self):
        for key, client in self.clients.items():
            missing = []
            for item in self.files_for_key.get(key, []):
                filename = item["name"]
                if ServiceDocument.objects.filter(
                    vehicle__client=client,
                    custom_name__iexact=filename[:150],
                ).exists():
                    self.skipped_documents += 1
                    continue
                stored_name = self._zip_index.get(filename.lower())
                if not stored_name:
                    missing.append(filename)
                    continue
                vehicle = None
                if item["vin"] or item["plate"]:
                    vehicle = self.vehicle_for(
                        client,
                        {"vin": item["vin"], "plate_number": item["plate"]},
                        create=False,
                    )
                if vehicle is None:
                    vehicle = client.vehicles.order_by("id").first()
                if vehicle is None:
                    missing.append(filename)
                    continue
                payload = self._zip.read(stored_name)
                document = ServiceDocument(
                    vehicle=vehicle,
                    document_type=document_type_for(filename),
                    custom_name=filename[:150],
                )
                safe_name = re.sub(r"[^A-Za-z0-9._-]+", "_", os.path.basename(filename))[:80] or "document"
                document.file.save(safe_name, ContentFile(payload), save=True)
                self.created_documents += 1
            if not missing:
                continue
            content = "Kintone files listed but not included in the zip: " + ", ".join(sorted(set(missing)))
            if not ClientNote.objects.filter(client=client, content=content).exists():
                ClientNote.objects.create(client=client, created_by=self.actor, content=content)

    def open_zip(self, uploaded):
        if uploaded is None:
            return
        uploaded.seek(0)
        self._zip = zipfile.ZipFile(uploaded)
        for info in self._zip.infolist():
            if info.is_dir():
                continue
            base = os.path.basename(info.filename)
            if not base or base.startswith("."):
                continue
            self._zip_index.setdefault(base.lower(), info.filename)

    def close(self):
        if self._zip is not None:
            self._zip.close()

    def result(self):
        return {
            "clients_created": self.created_clients,
            "clients_skipped": self.skipped_clients,
            "vehicles_created": self.created_vehicles,
            "vehicles_skipped": self.skipped_vehicles,
            "transactions_created": self.created_transactions,
            "transactions_skipped": self.skipped_transactions,
            "documents_created": self.created_documents,
            "documents_skipped": self.skipped_documents,
            "errors": self.errors[:40],
            "error_count": len(self.errors),
        }


def _consume_rows(importer, rows, kind):
    for index, row in enumerate(rows, start=2):
        key = normalize_key(row.get(importer.link_header))
        if not key:
            importer.errors.append(f"{kind} row {index} has no unique number, so it was skipped.")
            continue
        try:
            importer.note_files(key, row)
            client = importer.client_for(key, row)
            if client is None:
                continue
            if kind == "Vehicles":
                importer.vehicle_for(client, row, create=True)
            elif kind == "Transactions":
                importer.add_transaction(client, row)
        except Exception as exc:
            importer.errors.append(f"{kind} row {index}: {exc}")


def import_kintone(
    *,
    organization,
    actor,
    clients_file,
    vehicles_file=None,
    transactions_file=None,
    documents_file=None,
    documents_zip=None,
    link_column="",
):
    client_rows = rows_from_upload(clients_file)
    vehicle_rows = rows_from_upload(vehicles_file)
    transaction_rows = rows_from_upload(transactions_file)
    document_rows = rows_from_upload(documents_file)
    link_header = resolve_link_header(client_rows, link_column)
    require_link_header(vehicle_rows, link_header, "vehicles")
    require_link_header(transaction_rows, link_header, "transactions")
    require_link_header(document_rows, link_header, "documents")

    importer = KintoneImporter(organization, actor, link_header)
    try:
        importer.open_zip(documents_zip)
        _consume_rows(importer, client_rows, "Clients")
        _consume_rows(importer, vehicle_rows, "Vehicles")
        _consume_rows(importer, transaction_rows, "Transactions")
        _consume_rows(importer, document_rows, "Documents")
        importer.attach_files()
        return importer.result()
    finally:
        importer.close()
